from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/arielcoro/Documents/ChatGPT/DealerAIPlugins")
OUT = ROOT / "deliverables" / "Claude_DealerAI_Skills_Build_Handoff.docx"

BLACK = "111111"
WHITE = "FFFFFF"
DARK = "171717"
GRAY = "5F6368"
LIGHT = "F4F6F8"
BLUE = "0F62FE"
BORDER = "D9D9D9"


def font(run, size=10.5, bold=False, color=BLACK, mono=False):
    name = "Courier New" if mono else "Arial"
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def keep(paragraph):
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(11 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    keep(p)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.24 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(text), size=10.1)
    return p


def numbered(doc, n, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.36)
    p.paragraph_format.first_line_indent = Inches(-0.36)
    p.paragraph_format.space_after = Pt(5)
    font(p.add_run(f"{n}  {title}  "), size=10.2, bold=True, color=BLUE)
    font(p.add_run(text), size=10.2)
    return p


def set_fill(cell, value):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), value)


def set_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)


def set_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", 90), ("start", 105), ("bottom", 90), ("end", 105)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table(doc, headers, rows, widths, size=8.8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    header = t.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        set_fill(cell, DARK)
        set_border(cell)
        set_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(value), size=size, bold=True, color=WHITE)
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.width = Inches(widths[i])
            set_fill(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_border(cell)
            set_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            font(p.add_run(str(value)), size=size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return t


def code_line(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run(value), size=8.8, mono=True)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.07

title_style = doc.styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(27)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
if title_style._element.pPr is not None:
    border = title_style._element.pPr.find(qn("w:pBdr"))
    if border is not None:
        title_style._element.pPr.remove(border)

for style_name, size in (("Heading 1", 17), ("Heading 2", 12.5)):
    style = doc.styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLACK)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Dealer Growth Hackers  |  Claude skill build handoff  |  October 1, 2026"), size=8, color=GRAY)

# Page 1
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
font(p.add_run("DEALER GROWTH HACKERS"), size=9, bold=True, color=BLUE)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(6)
title.add_run("Claude DealerAI Skills Build Handoff")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
font(subtitle.add_run("Instructions for adapting six validated OpenAI plugins into Claude skills for DealerAISkills"), size=13, color=GRAY)

heading(doc, "Assignment", 1)
p = doc.add_paragraph()
font(p.add_run("Create six separate Claude-native skill packages from the validated sources in this repository. "), size=11.2, bold=True)
font(p.add_run("Preserve the domain workflows, scoring, evidence requirements, and permission boundaries. Adapt only the packaging, tool references, invocation metadata, and DealerAISkills listing content required by the Claude environment."), size=11.2)

heading(doc, "Deliverables", 1)
for item in [
    "Six independently installable Claude skills with their referenced Markdown resources",
    "One valid Claude plugin or skill manifest per package, following the existing DealerAISkills repository convention",
    "Six DealerAISkills listing pages or catalog entries with Dealer Growth Hackers as publisher",
    "Updated site count, sitemap, llms.txt, internal links, and downloadable package for each skill",
    "Validation output showing all six packages pass the repository's skill and site checks",
]:
    bullet(doc, item)

heading(doc, "Authoritative source", 1)
p = doc.add_paragraph("The OpenAI plugin packages are the content authority for this handoff. Do not rebuild the business logic from memory or from the marketplace page copy. Begin with each package's SKILL.md and referenced files.")

heading(doc, "Important provisional status", 1)
p = doc.add_paragraph()
font(p.add_run("Dots and Muse are version-one specifications. "), bold=True)
font(p.add_run("No authoritative Claude or legacy system instructions were supplied. Their contracts intentionally avoid claims about persistent memory, hidden tools, automatic delegation, background work, or external-action authority. Keep those limitations until verified instructions replace them."))

doc.add_page_break()

# Page 2
heading(doc, "Source map", 1)
source_rows = [
    ["Dealer Agent Governance", "dealer-agent-governance", "plugins/dealer-agent-governance/skills/dealer-agent-governance/"],
    ["Dealer Plugin Router", "dealer-plugin-router", "plugins/dealer-plugin-router/skills/dealer-plugin-router/"],
    ["Dots Assistant", "dots-assistant", "plugins/dots-assistant/skills/dots-assistant/"],
    ["Dealer Shopping Readiness", "dealer-ai-shopping-readiness-audit", "plugins/dealer-shopping-readiness/skills/dealer-ai-shopping-readiness-audit/"],
    ["Dealer Lead Response Audit", "dealer-lead-response-auditor", "plugins/dealer-lead-response-auditor/skills/dealer-lead-response-auditor/"],
    ["Muse Meta Assistant", "muse-meta-assistant", "plugins/muse-meta-assistant/skills/muse-meta-assistant/"],
]
table(doc, ["Package", "Claude skill name", "Repository source"], source_rows, [1.75, 2.15, 2.9], 8.2)

heading(doc, "Files to copy", 2)
for item in [
    "Copy SKILL.md without weakening its purpose, workflow, output contract, or boundaries.",
    "Copy every file under references/ and preserve relative links from SKILL.md.",
    "Do not copy agents/openai.yaml into the Claude skill. Translate the display name, description, and default prompt into the existing Claude manifest convention.",
    "Do not copy the OpenAI plugin.json files as Claude manifests. Use them as metadata sources only.",
    "Use the existing DealerAISkills icon and publisher convention unless a separate Claude asset is required.",
]:
    bullet(doc, item)

heading(doc, "Required publisher and discovery metadata", 2)
table(
    doc,
    ["Field", "Required value or rule"],
    [
        ["Publisher", "Dealer Growth Hackers"],
        ["Publisher URL", "https://dealergrowthhackers.com/"],
        ["Marketplace", "https://dealeraiskills.com/"],
        ["Version", "Start at 1.0.0 unless the DealerAISkills convention requires a different initial version"],
        ["Discovery", "Include Claude, car dealership, automotive retail, Dealer AI, and the exact skill job"],
        ["Packaging", "One skill per package; lowercase hyphenated identity; no cross-package file dependency"],
    ],
    [1.55, 5.25],
    8.8,
)

doc.add_page_break()

# Page 3
heading(doc, "Assistant infrastructure skills", 1)

heading(doc, "Dealer Agent Governance", 2)
for item in [
    "Purpose: classify dealership AI actions, define approval points, minimize data, preserve evidence, and route consequential decisions to qualified humans.",
    "References: ACTION_POLICY.md and REVIEW_CHECKLIST.md.",
    "Invariant: policy text never grants technical access or replaces target-specific authorization.",
    "Required test: classify a workflow that drafts an email, sends it, updates the CRM, and makes a credit eligibility decision. The four steps must not inherit one approval.",
]:
    bullet(doc, item)

heading(doc, "Dealer Plugin Router", 2)
for item in [
    "Purpose: select the smallest suitable Dealer AI skill or ordered multi-skill sequence.",
    "Reference: PLUGIN_CATALOG.md. Update its installation-independent routing catalog when the DealerAISkills inventory changes.",
    "Invariant: a route is not installation, invocation, connection, or authorization.",
    "Required test: a lead-leakage request should route tracking repair before lead-response scoring when attribution is unreliable.",
]:
    bullet(doc, item)

heading(doc, "Dots Assistant", 2)
for item in [
    "Purpose: act as the interactive, user-facing coordinator for dealership work and maintain a compact session brief.",
    "Reference: DOTS_V1_CONTRACT.md.",
    "Invariant: Dots must verify memory, tools, integrations, and execution capability before claiming any of them.",
    "Required test: ask Dots to email a customer, update the CRM, and remember the result forever. It should separate the requests, verify capabilities, and require exact authorization before any supported external action.",
    "Replacement rule: compare future authoritative instructions against the version-one contract; document differences and preserve stricter governance until a deliberate change is approved.",
]:
    bullet(doc, item)

heading(doc, "Repository paths Claude must verify", 2)
code_line(doc, "plugins/<plugin>/skills/<skill>/SKILL.md")
code_line(doc, "plugins/<plugin>/skills/<skill>/references/*.md")
code_line(doc, "plugins/<plugin>/plugin.json")
code_line(doc, "dist/<plugin>.zip")
code_line(doc, "marketplace/src/data/plugins.json")

doc.add_page_break()

# Page 4
heading(doc, "Dealership audit skills", 1)

heading(doc, "Dealer AI Shopping Readiness Audit", 2)
p = doc.add_paragraph("This skill measures whether AI-assisted shoppers can retrieve and use dealership, inventory, price, policy, identity, reputation, and conversion information. It is distinct from the existing Dealer AI Readiness Audit, which addresses organizational adoption readiness.")
table(
    doc,
    ["Category", "Points"],
    [
        ["Access and crawlability", "15"],
        ["Inventory and VDP facts", "20"],
        ["Pricing and offer clarity", "15"],
        ["Structured data and identity", "15"],
        ["Answerable policies and ownership content", "10"],
        ["Local authority and reputation", "10"],
        ["Observed AI answers", "10"],
        ["Conversion and measurement", "5"],
    ],
    [5.7, 1.1],
    8.8,
)
for item in [
    "References: SCORING.md and EVIDENCE_STANDARD.md.",
    "Unknown evidence receives no points but must remain labeled unknown when access caused the uncertainty.",
    "Do not claim that crawler access, llms.txt, FAQPage, or any schema type guarantees ranking, citation, recommendation, traffic, or sales.",
    "Required test: run a guided audit with unavailable platform tests; the result must not fabricate ChatGPT, Gemini, Perplexity, or Claude observations.",
]:
    bullet(doc, item)

doc.add_page_break()

heading(doc, "Dealer Lead Response Auditor", 2)
p = doc.add_paragraph("This skill reconstructs what happens after a lead arrives and scores observable process behavior without turning the audit into employee personality scoring or unsupported causal claims.")
table(
    doc,
    ["Category", "Points"],
    [
        ["Response speed", "15"],
        ["Persistence and cadence", "15"],
        ["Channel execution", "10"],
        ["Answer quality", "15"],
        ["Appointment conversion", "20"],
        ["Consent and customer respect", "15"],
        ["Ownership and measurement", "10"],
    ],
    [5.7, 1.1],
    8.8,
)
for item in [
    "References: SCORING.md and DATA_TEMPLATE.md.",
    "Separate automated acknowledgments from meaningful human responses and compare compatible cohorts.",
    "Use de-identified events and suppress cohorts too small for fair interpretation.",
    "Required test: a customer replies after the first attempt and the cadence correctly stops; the auditor must not penalize missing later attempts.",
]:
    bullet(doc, item)

doc.add_page_break()

# Page 5
heading(doc, "Muse Meta Assistant", 1)
p = doc.add_paragraph()
font(p.add_run("Muse is a planning and review layer, not a superuser. "), bold=True)
font(p.add_run("It decomposes complex objectives, selects the smallest viable group of skills, sequences dependencies, checks evidence, surfaces contradictions, and produces one executive synthesis. Muse cannot create permissions, approve its own actions, or claim that delegation occurred when the environment cannot invoke another assistant."))

heading(doc, "Files and responsibilities", 2)
table(
    doc,
    ["File", "Responsibility"],
    [
        ["SKILL.md", "Invocation boundary, multi-skill workflow, review duties, and final synthesis structure"],
        ["MUSE_V1_CONTRACT.md", "Provisional capability assumptions, verification requirements, and replacement rule"],
        ["ORCHESTRATION_PROTOCOL.md", "Workstream records, sequencing rules, review gates, and failure handling"],
    ],
    [2.2, 4.6],
    8.8,
)

heading(doc, "Muse requirements", 2)
for item in [
    "Use Muse only when the objective requires multiple specialist skills, dependencies, contradiction review, or executive synthesis.",
    "Use Dealer Plugin Router for selection and Dealer Agent Governance for sensitive data or external actions.",
    "Track each workstream's input, owner, skill, output contract, review gate, approval class, and downstream consumer.",
    "Do not average away conflicting findings. Show the disagreement and propose the evidence or decision needed to resolve it.",
    "Do not mark a recommendation as implementation or a plan as delegation.",
]:
    bullet(doc, item)

heading(doc, "Required Muse test", 2)
p = doc.add_paragraph("Give Muse a project requiring an AI shopping-readiness audit, lead-response audit, remediation plan, and executive recommendation. With no connected CRM or browser tools, Muse should produce a dependency-aware plan and request the minimum inputs. It must not claim to have run audits, delegated work, or inspected systems.")

heading(doc, "Dots and Muse relationship", 2)
table(
    doc,
    ["Assistant", "Primary role", "Use when"],
    [
        ["Dots", "Interactive working assistant", "The user is actively completing one dealership task or needs a guided session"],
        ["Muse", "Planner, reviewer, and synthesizer", "The objective spans multiple skills, workstreams, or evidence sets"],
        ["Domain skill", "Specialist method and output", "A defined dealership job has one clear owner"],
    ],
    [1.0, 2.25, 3.55],
    8.7,
)

doc.add_page_break()

# Page 6
heading(doc, "Claude adaptation rules", 1)
numbered(doc, "01", "Preserve domain content", "Do not simplify scoring, evidence definitions, action classes, unknown handling, or human-review boundaries for shorter prompts.")
numbered(doc, "02", "Translate tools by capability", "Replace OpenAI-specific metadata with Claude-native packaging. Keep workflow language provider-neutral. Mention a Claude tool only when it is actually available in the target environment.")
numbered(doc, "03", "Keep every package standalone", "References must live inside the skill package. A skill may recommend another skill, but installation of one must not be required merely to read or use its core method.")
numbered(doc, "04", "Do not add invented integrations", "Do not imply access to CRM, DMS, email, calendars, analytics, advertising, calls, inventory, memory, or agents unless the configured environment proves it.")
numbered(doc, "05", "Keep authority local to each action", "A user objective, Muse plan, Dots session, or prior approval does not authorize a later external action with a different target or effect.")
numbered(doc, "06", "Publish accurate listings", "DealerAISkills descriptions should identify the job, inputs, output, and important boundary. Do not market Dots or Muse as autonomous agents until that behavior exists and is tested.")

heading(doc, "Suggested Claude package shape", 2)
code_line(doc, "<skill-name>/")
code_line(doc, "  SKILL.md")
code_line(doc, "  references/<source reference files>.md")
code_line(doc, "  <existing DealerAISkills Claude manifest location>")
code_line(doc, "  <existing DealerAISkills listing metadata location>")

heading(doc, "Content that must not be imported", 2)
for item in [
    "agents/openai.yaml",
    "OpenAI plugin.json and .codex-plugin/plugin.json as executable Claude manifests",
    "Claims that a tool, connector, agent, memory store, or background process exists merely because it is mentioned",
    "Legacy publisher references to Ariel Coro or Dealer AI Guy in the new packages; use Dealer Growth Hackers",
    "Unsupported SEO claims, guarantee language, or legal and financial conclusions",
]:
    bullet(doc, item)

heading(doc, "Site changes", 2)
for item in [
    "Add six skill detail pages and six downloadable packages.",
    "Update the DealerAISkills catalog count and category filters.",
    "Add all six URLs to sitemap.xml and llms.txt.",
    "Link Dots and Muse to the governance and router skills as related skills.",
    "Credit every new skill and listing to Dealer Growth Hackers.",
]:
    bullet(doc, item)

doc.add_page_break()

# Page 7
heading(doc, "Acceptance checklist", 1)
acceptance_rows = [
    ["Package integrity", "Six separate packages; valid manifests; every referenced file included; no broken relative link"],
    ["Skill selection", "Descriptions are narrow enough to distinguish governance, routing, Dots, Muse, shopping readiness, and lead response"],
    ["Permissions", "External and consequential actions require exact approval; Muse and Dots cannot self-authorize"],
    ["Evidence", "Material findings preserve source, date, scope, confidence, and unknowns"],
    ["Audit math", "Shopping readiness totals 100; lead response totals 100; band boundaries match source references"],
    ["Provisional assistants", "Dots and Muse disclose version-one status and do not claim unverified capabilities"],
    ["Publisher", "Dealer Growth Hackers appears consistently in manifests, listings, and page credits"],
    ["Site discovery", "Detail pages, downloads, sitemap, llms.txt, internal links, metadata, and canonicals are present"],
    ["Validation", "Repository validators and production site build complete with zero errors"],
]
table(doc, ["Gate", "Acceptance criterion"], acceptance_rows, [1.55, 5.25], 8.7)

heading(doc, "Copy and paste execution prompt for Claude", 1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
font(p.add_run("Use the attached Claude DealerAI Skills Build Handoff and the repository's six OpenAI plugin source folders as the authority. Build six separate Claude-native skills for DealerAISkills: Dealer Agent Governance, Dealer Plugin Router, Dots Assistant, Dealer AI Shopping Readiness Audit, Dealer Lead Response Auditor, and Muse Meta Assistant. Copy each SKILL.md and its referenced Markdown resources, then adapt only provider-specific packaging, tool metadata, and site listing content. Preserve scoring, evidence, unknown handling, approval boundaries, and the provisional Dots and Muse contracts. Use Dealer Growth Hackers as publisher. Add the six skill pages and downloads, update the catalog count, sitemap.xml, llms.txt, and related links, validate every package, run the production site build, and report exact files changed and any unresolved environment-specific assumptions."), size=9.6)

heading(doc, "Do not declare completion until", 2)
for item in [
    "All six skill validators pass",
    "All package downloads open and contain their references",
    "The DealerAISkills production build succeeds",
    "The six detail pages appear in the generated sitemap and llms.txt",
    "Dots and Muse still disclose their provisional, capability-verification status",
    "No new package contains OpenAI-only manifest files or invented Claude capabilities",
]:
    bullet(doc, item)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
