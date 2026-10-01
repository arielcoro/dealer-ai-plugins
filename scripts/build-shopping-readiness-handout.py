from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/arielcoro/Documents/ChatGPT/DealerAIPlugins")
OUT = ROOT / "deliverables" / "Dealer_AI_Shopping_Readiness_Audit_Website_Handout.docx"

BLACK = "111111"
WHITE = "FFFFFF"
DARK = "171717"
GRAY = "5F6368"
LIGHT = "F3F4F6"
LIGHT_BLUE = "EEF4FF"
BLUE = "0F62FE"
BORDER = "D9D9D9"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_font(run, name="Arial", size=10.5, bold=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    r_pr.append(r_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.23 + level * 0.2)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    r = p.add_run(text)
    set_font(r, size=9.7)
    return p


def add_number(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.32)
    r = p.add_run(f"{number}  {title}  ")
    set_font(r, size=10, bold=True, color=BLUE)
    r2 = p.add_run(text)
    set_font(r2, size=10)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_fill(cell, DARK)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=font_size, bold=True, color=WHITE)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_fill(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text))
            set_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(28)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
if title_style._element.pPr is not None:
    title_border = title_style._element.pPr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style._element.pPr.remove(title_border)

for style_name, size in (("Heading 1", 17), ("Heading 2", 12.5)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLACK)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Dealer Growth Hackers  |  Website implementation handout  |  October 1, 2026")
set_font(fr, size=8, color=GRAY)

# Page 1
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("DEALER GROWTH HACKERS")
set_font(r, size=9, bold=True, color=BLUE)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(6)
title.add_run("Dealer AI Shopping Readiness Audit")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(15)
r = subtitle.add_run("Website additions required to launch the audit as a search and lead-generation offer")
set_font(r, size=13, color=GRAY)

add_heading(doc, "Recommendation", 1)
p = doc.add_paragraph()
r = p.add_run("Add one dedicated audit landing page, one supporting AI car-buying content hub, and a short conversion workflow. ")
set_font(r, size=11, bold=True)
r = p.add_run("The landing page should explain how AI shopping assistants interpret a dealership's inventory, pricing, policies, reputation, and website. It should generate qualified audit requests while reinforcing Dealer Growth Hackers' existing promise that one team owns paid media, SEO, site, and conversion.")
set_font(r, size=11)

add_heading(doc, "Search demand supporting the offer", 1)
p = doc.add_paragraph("The AnswerThePublic export contains 38 unique queries. Four have reported monthly volume. The root phrase has the highest commercial signal in the file.")
p.paragraph_format.space_after = Pt(7)
add_table(
    doc,
    ["Search query", "Reported volume", "CPC"],
    [
        ["ai car buying", "50", "$12.36"],
        ["ai car buying app", "30", "$2.94"],
        ["ai car buying assistant", "20", "$4.30"],
        ["ai for car buying", "10", "$3.83"],
    ],
    widths=[4.3, 1.35, 1.15],
    font_size=9.2,
)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(2)
r = p.add_run("Interpretation  ")
set_font(r, size=9.5, bold=True)
r = p.add_run("Demand is early and concentrated. Build one strong offer and one supporting topic cluster rather than several overlapping products.")
set_font(r, size=9.5)

add_heading(doc, "Pages to add", 1)
add_number(doc, "01", "Primary audit page", "/dealer-ai-shopping-readiness-audit/ converts search and referral traffic into audit requests.")
add_number(doc, "02", "AI car-buying hub", "/ai-car-buying/ explains the shift in buyer behavior and links to the audit.")
add_number(doc, "03", "Supporting articles", "Publish focused pages for AI car-buying assistants, prompts, inventory interpretation, and dealership readiness.")
add_number(doc, "04", "Marketplace connection", "Link the future Dealer AI Shopping Readiness Audit plugin from dealeraiplugins.com after the plugin is built and published.")

doc.add_page_break()

# Page 2
add_heading(doc, "Primary audit landing page", 1)
p = doc.add_paragraph()
r = p.add_run("Recommended URL  ")
set_font(r, bold=True)
r = p.add_run("https://dealergrowthhackers.com/dealer-ai-shopping-readiness-audit/")
set_font(r, color=BLUE)

add_heading(doc, "Search metadata and hero copy", 2)
add_table(
    doc,
    ["Element", "Recommended copy"],
    [
        ["SEO title", "Dealer AI Shopping Readiness Audit | Dealer Growth Hackers"],
        ["Meta description", "Find out whether AI car-buying assistants can understand your inventory, pricing, policies, reputation, and dealership website."],
        ["H1", "Can AI car-buying assistants understand and recommend your inventory"],
        ["Hero text", "We test what AI shopping assistants can find, understand, and explain about your dealership, then show you which website and data gaps need attention."],
        ["Primary CTA", "Request the audit"],
        ["Secondary CTA", "See what the audit covers"],
    ],
    widths=[1.35, 5.45],
    font_size=8.9,
)

add_heading(doc, "Required page sections", 2)
sections = [
    ("1", "Buyer behavior", "Explain that shoppers increasingly use ChatGPT and other AI assistants to compare vehicles, ownership costs, dealerships, policies, and offers."),
    ("2", "Why dealers disappear", "Show the practical causes: thin VDP data, inconsistent pricing, blocked crawlers, unclear policies, weak entity signals, and incomplete structured data."),
    ("3", "What the audit checks", "Present the eight audit categories and the evidence reviewed in each category."),
    ("4", "What the dealer receives", "List the readiness score, evidence appendix, prioritized fixes, sample AI responses, and 30-day action plan."),
    ("5", "How it works", "Request, evidence collection, testing, findings review, and implementation planning."),
    ("6", "Who should use it", "Dealer principals, general managers, marketing leaders, dealer groups, and agencies responsible for website performance."),
    ("7", "Proof and methodology", "Explain the test prompts, sources, date of testing, limitations, and the difference between observed results and recommendations."),
    ("8", "Conversion form", "Place a short form near the hero and repeat it after the deliverables section."),
]
for n, heading, text in sections:
    add_number(doc, n.zfill(2), heading, text)

add_heading(doc, "Lead form fields", 2)
for item in [
    "Dealership or dealer group name",
    "Website URL and primary market",
    "New, used, or both",
    "Website platform or provider, if known",
    "Name, role, business email, and phone",
    "Consent checkbox and link to the privacy policy",
]:
    add_bullet(doc, item)

doc.add_page_break()

# Page 3
add_heading(doc, "Recommended audit framework", 1)
p = doc.add_paragraph("Use a 100-point score only after the checks and evidence standards are finalized. The weights below provide a practical starting structure for the website and the future plugin.")
p.paragraph_format.space_after = Pt(8)

audit_rows = [
    ["Crawl and access", "15", "robots.txt, sitemap, status codes, canonical URLs, important content in initial HTML, access for search and user-triggered AI fetchers"],
    ["Inventory and VDP data", "20", "VIN, price, mileage, trim, availability, photos, descriptions, options, incentives, freshness, and consistent inventory feeds"],
    ["Pricing and offer clarity", "15", "price conditions, fees, incentives, financing assumptions, trade-in language, expiration dates, and disclosure visibility"],
    ["Structured data and entities", "15", "Vehicle, Product, Offer, AutoDealer or LocalBusiness, dealership identity, NAP consistency, and valid page relationships"],
    ["Answerable dealership content", "10", "service, finance, returns, delivery, warranties, trade-ins, appointment policies, and direct answers to buyer questions"],
    ["Local authority and reputation", "10", "Google Business Profile, review patterns, third-party references, staff and location details, and consistent brand facts"],
    ["AI answer testing", "10", "buyer-intent prompt set, accuracy, source attribution, omissions, contradictions, competitor comparisons, and dated evidence"],
    ["Conversion and measurement", "5", "clear next actions, working forms and phone links, GA4 events, source attribution, and referral tracking"],
]
add_table(doc, ["Audit category", "Weight", "Evidence reviewed"], audit_rows, widths=[1.65, 0.65, 4.5], font_size=8.3)

add_heading(doc, "Dealer deliverables", 2)
for item in [
    "Executive readiness score with section-level findings",
    "Evidence appendix with tested URLs, prompts, dates, and observed AI responses",
    "Critical, high, medium, and low-priority remediation list",
    "Thirty-day implementation plan with owners and acceptance criteria",
    "Measurement plan for organic search, AI referrals, form starts, submissions, and qualified opportunities",
]:
    add_bullet(doc, item)

add_heading(doc, "Methodology language required on the page", 2)
p = doc.add_paragraph("State that AI answers vary by platform, location, session, and date. The audit reports observable gaps and recommended fixes; it does not guarantee rankings, citations, traffic, or sales. Do not claim that llms.txt, FAQPage markup, or any single schema type causes AI recommendations.")

doc.add_page_break()

# Page 4
add_heading(doc, "Supporting content and internal links", 1)
content_rows = [
    ["/ai-car-buying/", "AI Car Buying Guide for Dealerships", "Pillar page targeting the root topic and explaining the operational implications for dealers."],
    ["/ai-car-buying-assistant/", "How AI Car-Buying Assistants Evaluate Dealers", "Explains the information assistants need and links to the audit."],
    ["/ai-car-buying-prompts/", "AI Car-Buying Prompts and What They Reveal", "Uses realistic buyer prompts to show gaps dealers should test."],
    ["/ai-car-buying-inventory-data/", "Inventory Data AI Assistants Can Understand", "Covers VDP fields, freshness, feeds, and structured data."],
    ["/dealer-ai-shopping-readiness-audit/", "Dealer AI Shopping Readiness Audit", "Primary conversion page and canonical explanation of the service."],
]
add_table(doc, ["URL", "Page title", "Purpose"], content_rows, widths=[2.1, 2.35, 2.35], font_size=8.4)

p = doc.add_paragraph()
r = p.add_run("Internal links  ")
set_font(r, bold=True)
r = p.add_run("Link the audit from the homepage, Growth Consulting page, About page, relevant AI-search articles, and the site footer. Link every supporting article back to the audit page with descriptive anchor text.")
set_font(r)

add_heading(doc, "Technical requirements", 1)
for item in [
    "Prerender or server-render the audit and supporting pages. The current homepage source contains an empty root element and relies on client-side JavaScript for visible content.",
    "Return a 200 status with the title, meta description, H1, body copy, canonical URL, and structured data in the initial HTML.",
    "Add the pages to the XML sitemap and confirm robots.txt does not block them.",
    "Use Organization, WebSite, WebPage or Service, and BreadcrumbList structured data where the visible page supports it.",
    "Create one self-referencing canonical URL per page and redirect alternate slash, www, and parameter versions consistently.",
    "Keep the lead form usable without layout shifts and label every field for accessibility.",
    "Add Open Graph and social preview metadata with a branded audit image.",
]:
    add_bullet(doc, item)

add_heading(doc, "Measurement plan", 1)
add_table(
    doc,
    ["Event or metric", "Implementation"],
    [
        ["Audit CTA click", "GA4 event with page path and CTA location"],
        ["Form start", "Fire after the first intentional field interaction"],
        ["Form submit", "Fire only after confirmed successful submission"],
        ["Qualified opportunity", "Import CRM stage or offline conversion with source preserved"],
        ["Organic demand", "Track impressions, clicks, average position, and landing page by query cluster in Search Console"],
        ["AI referrals", "Report known AI referrers separately and retain landing page and conversion outcome"],
    ],
    widths=[1.65, 5.15],
    font_size=8.7,
)

doc.add_page_break()

# Page 5
add_heading(doc, "Launch checklist", 1)
launch_rows = [
    ["Offer", "Finalize scope, evidence standards, score weights, turnaround, and whether pricing appears publicly.", "Leadership"],
    ["Copy", "Approve landing-page copy, methodology limits, deliverables, and CTA language.", "Marketing"],
    ["Design", "Create responsive page modules and a branded social preview image.", "Design"],
    ["Development", "Build prerendered routes, accessible form, structured data, redirects, and sitemap entries.", "Web"],
    ["Analytics", "Configure GA4 events, CRM attribution, and Search Console reporting.", "Analytics"],
    ["Audit operations", "Create the evidence template, prompt set, report template, and quality-review checklist.", "Delivery"],
    ["Plugin", "Build and validate the Dealer AI Shopping Readiness Audit plugin, then link its public listing.", "Product"],
    ["QA", "Test mobile, form delivery, initial HTML, metadata, schema, analytics, and privacy language.", "Web and Marketing"],
]
add_table(doc, ["Workstream", "Acceptance criterion", "Owner"], launch_rows, widths=[1.1, 4.85, 0.85], font_size=8.5)

add_heading(doc, "Recommended launch order", 1)
add_number(doc, "01", "Define the offer", "Finalize the audit scope and evidence standards before publishing sales copy.")
add_number(doc, "02", "Build the landing page", "Launch the prerendered audit page, form, analytics, and core internal links.")
add_number(doc, "03", "Publish the pillar page", "Add the AI car-buying guide and link it to the audit.")
add_number(doc, "04", "Release supporting content", "Publish one focused article at a time and measure qualified traffic rather than raw pageviews.")
add_number(doc, "05", "Publish the plugin", "Add the public ChatGPT and Codex directory listing after review and connect it to both sites.")

add_heading(doc, "Sources", 1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Search-demand source  ")
set_font(r, size=9, bold=True)
r = p.add_run("AnswerThePublic export for “ai car buying,” United States, English, October 1, 2026.")
set_font(r, size=9)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Current website reviewed  ")
set_font(r, size=9, bold=True)
add_hyperlink(p, "dealergrowthhackers.com", "https://dealergrowthhackers.com/")

p = doc.add_paragraph()
r = p.add_run("Plugin marketplace  ")
set_font(r, size=9, bold=True)
add_hyperlink(p, "dealeraiplugins.com", "https://dealeraiplugins.com/")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
