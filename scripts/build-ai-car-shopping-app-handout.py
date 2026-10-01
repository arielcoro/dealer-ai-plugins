from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/arielcoro/Documents/ChatGPT/DealerAIPlugins")
OUT = ROOT / "deliverables" / "AI_Assisted_Car_Shopping_App_Handout.docx"

BLACK = "111111"
WHITE = "FFFFFF"
DARK = "171717"
GRAY = "5F6368"
MID = "D9D9D9"
LIGHT = "F5F7FA"
LIGHT_BLUE = "EDF4FF"
BLUE = "0F62FE"
GREEN = "198038"
AMBER = "B28600"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=MID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=10.2, bold=False, color=BLACK):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def add_hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(11 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    keep_with_next(p)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.23 + 0.18 * level)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    set_font(p.add_run(text), size=9.5)
    return p


def add_step(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.34)
    p.paragraph_format.first_line_indent = Inches(-0.34)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(f"{number}  {title}  "), size=9.8, bold=True, color=BLUE)
    set_font(p.add_run(text), size=9.8)
    return p


def add_label_paragraph(doc, label, text, color=BLACK):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(label + "  "), size=9.8, bold=True, color=color)
    set_font(p.add_run(text), size=9.8)
    return p


def add_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        set_cell_fill(cell, DARK)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), size=font_size, bold=True, color=WHITE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.width = Inches(widths[i])
            set_cell_fill(cell, WHITE if row_index % 2 == 0 else LIGHT)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(text)), size=font_size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.56)
section.bottom_margin = Inches(0.56)
section.left_margin = Inches(0.66)
section.right_margin = Inches(0.66)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.06

title_style = doc.styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(27)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
if title_style._element.pPr is not None:
    border = title_style._element.pPr.find(qn("w:pBdr"))
    if border is not None:
        title_style._element.pPr.remove(border)

for style_name, size in (("Heading 1", 17), ("Heading 2", 12.2)):
    style = doc.styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLACK)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Dealer Growth Hackers  |  Product concept handout  |  October 1, 2026"), size=8, color=GRAY)

# Page 1
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run("DEALER GROWTH HACKERS"), size=9, bold=True, color=BLUE)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(6)
title.add_run("AI Assisted Car Shopping App")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
set_font(subtitle.add_run("Product concept, launch scope, trust model, and domain recommendation"), size=13, color=GRAY)

add_heading(doc, "Executive recommendation", 1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(9)
set_font(p.add_run("Build a consumer decision assistant—not another vehicle-results page. "), size=11.2, bold=True)
set_font(p.add_run("The app should help a shopper turn needs, budget, listings, quotes, and financing assumptions into a clear shortlist and an evidence-backed buying plan. It should show sources and dates, expose uncertainty, and make referral economics visible."), size=11.2)

add_heading(doc, "Recommended working brand", 1)
add_table(
    doc,
    ["Decision", "Recommendation", "Why"],
    [["Working name", "Better Car Buyer", "Outcome-led, credible, and broader than inventory search"],
     ["Primary domain", "BetterCarBuyer.com", "No .com registration record observed on Oct. 1, 2026"],
     ["AI-forward backup", "MyNextCarAI.com", "Immediately communicates the category"],
     ["Brand relationship", "Independent consumer brand, endorsed by Dealer Growth Hackers", "Separates consumer trust from dealer-service marketing"]],
    [1.25, 2.15, 3.4],
    8.8,
)

add_heading(doc, "The job to be done", 1)
p = doc.add_paragraph("When I am considering a vehicle, help me understand what fits my life and budget, compare the real tradeoffs, verify the deal inputs, and walk into the dealership knowing what to ask and what not to accept.")
p.paragraph_format.space_after = Pt(8)

add_heading(doc, "Search signal", 1)
add_table(
    doc,
    ["Query", "Reported monthly volume", "CPC"],
    [["ai car buying", "50", "$12.36"], ["ai car buying app", "30", "$2.94"], ["ai car buying assistant", "20", "$4.30"], ["ai for car buying", "10", "$3.83"]],
    [3.8, 1.75, 1.25],
    8.9,
)
add_label_paragraph(doc, "Interpretation", "This is an early category with clear commercial intent. Win with trust, direct answers, and a narrow valuable workflow before expanding into an inventory marketplace.")

doc.add_page_break()

# Page 2
add_heading(doc, "Audience and positioning", 1)
add_heading(doc, "Primary users", 2)
for item in [
    "First-time or infrequent buyers who do not know the process or vocabulary",
    "Budget-conscious shoppers comparing payment, cash price, financing, and ownership cost",
    "Families and commuters choosing among several vehicle types or powertrains",
    "Used-car shoppers who need a disciplined listing, history, inspection, and offer checklist",
    "Shoppers who already found a vehicle and want an independent second look",
]:
    add_bullet(doc, item)

add_heading(doc, "Positioning statement", 2)
p = doc.add_paragraph()
set_font(p.add_run("For car shoppers who want confidence without becoming automotive experts, "), bold=True)
set_font(p.add_run("Better Car Buyer is an AI-assisted decision guide that turns preferences, listings, and dealer quotes into a sourced comparison and step-by-step buying plan. Unlike generic chatbots or listing sites, it separates facts from estimates, dates every market-dependent claim, and explains how it gets paid."))

add_heading(doc, "Message hierarchy", 2)
add_table(
    doc,
    ["Message", "Proof the product must provide"],
    [["Choose with confidence", "A transparent fit score with user-controlled priorities"],
     ["See the real tradeoffs", "Purchase price, payment assumptions, ownership cost, safety, utility, and depreciation separated"],
     ["Understand the offer", "Itemized out-the-door comparison with fees and add-ons called out"],
     ["Know what to verify", "Source, date, uncertainty, and a human-verification step for every consequential claim"],
     ["Stay in control", "No dealer contact, credit action, or data sharing without explicit consent"]],
    [2.15, 4.65],
    8.7,
)

add_heading(doc, "Do not promise", 2)
for item in [
    "The guaranteed lowest price, guaranteed savings, or a guaranteed vehicle outcome",
    "Real-time price, availability, incentives, recalls, history, or financing without a live verified source",
    "That an AI recommendation replaces an inspection, test drive, mechanic, lender, insurer, or legal advice",
    "Unbiased recommendations if referral compensation or preferred inventory is not clearly disclosed",
]:
    add_bullet(doc, item)

doc.add_page_break()

# Page 3
add_heading(doc, "Minimum viable product", 1)
add_heading(doc, "The seven-step shopping journey", 2)
journey = [
    ("01", "Define the mission", "Capture use case, location, passengers, driving pattern, must-haves, timeline, cash or finance preference, and comfortable total budget."),
    ("02", "Build the shortlist", "Recommend three to five models with explicit reasons, compromises, and questions that could change the ranking."),
    ("03", "Compare real listings", "Accept URLs, screenshots, PDFs, or pasted details; normalize trim, mileage, price, fees, and seller claims."),
    ("04", "Model the cost", "Separate price, taxes, fees, down payment, APR, term, insurance estimate, energy or fuel, maintenance, and depreciation assumptions."),
    ("05", "Review the offer", "Extract an itemized dealer quote, flag unsupported add-ons or missing inputs, and compare offers on the same basis."),
    ("06", "Prepare to verify", "Generate test-drive, inspection, history, recall, warranty, financing, and paperwork questions."),
    ("07", "Create the decision brief", "Export a shareable summary of options, evidence, open risks, next actions, and the shopper's chosen priorities."),
]
for row in journey:
    add_step(doc, *row)

add_heading(doc, "MVP feature boundary", 2)
add_table(
    doc,
    ["Ship first", "Add after validation"],
    [["Conversational needs and budget intake", "Live nationwide inventory ingestion"],
     ["Model shortlist with explainable scoring", "Automated dealer outreach and negotiation"],
     ["Listing and quote upload or paste", "Credit prequalification and lender marketplace"],
     ["Side-by-side vehicle and offer comparison", "Trade-in bidding and transaction workflow"],
     ["Scenario calculator and verification checklist", "Native iOS and Android applications"],
     ["Cited decision brief and saved project", "White-label dealer, lender, and credit-union editions"]],
    [3.4, 3.4],
    8.8,
)

add_label_paragraph(doc, "MVP format", "Launch as a responsive web app. A web product is faster to test, easier to index, and better suited to shared listing links and uploaded quotes. Build native apps only after retention and repeat usage justify them.")

doc.add_page_break()

# Page 4
add_heading(doc, "Trust, data, and compliance", 1)
add_heading(doc, "Trust architecture", 2)
trust_rows = [
    ["Source layer", "Link to the inventory, manufacturer, government, lender, or user-provided document supporting each claim."],
    ["Time layer", "Show when price, availability, incentive, recall, rate, and listing facts were checked."],
    ["Confidence layer", "Label verified facts, calculated estimates, user inputs, AI inferences, and unknowns differently."],
    ["Consent layer", "Require a separate, specific approval before contacting a dealer, sharing identity, or starting a credit-related action."],
    ["Commercial layer", "Disclose referral, affiliate, preferred-placement, and lead fees beside the affected recommendation or outbound action."],
    ["Human layer", "Escalate high-stakes questions and encourage independent inspections and review of final documents."],
]
add_table(doc, ["Layer", "Product requirement"], trust_rows, [1.35, 5.45], 8.7)

add_heading(doc, "Compliance review before launch", 2)
for item in [
    "Advertising and referral disclosures: review FTC endorsement, native-advertising, and lead-generation requirements.",
    "Communications consent: obtain and retain channel-specific consent before calls or texts; review TCPA and state rules.",
    "Privacy: publish data-use, retention, deletion, sharing, and model-training policies; review applicable state privacy laws.",
    "Financing: do not imply approval or final terms. Credit, prequalification, or eligibility features need specialized FCRA, ECOA, fair-lending, and lender review.",
    "Recommendations: test for proxy discrimination and do not personalize access, price, or financing based on protected characteristics.",
    "Vehicle condition and safety: never convert an unverified seller claim into a product fact; recalls and history require authoritative sources and dates.",
]:
    add_bullet(doc, item)

add_heading(doc, "Business model guardrails", 2)
add_table(
    doc,
    ["Model", "Role", "Guardrail"],
    [["Free", "Needs intake, basic shortlist, and one comparison", "Useful without surrendering contact details"],
     ["Premium", "Unlimited projects, quote review, richer cost modeling, export", "Consumer pays for deeper help; no savings guarantee"],
     ["Concierge", "Optional human review or negotiation support", "Flat, published fee and defined scope"],
     ["Referral", "Dealer, lender, inspection, warranty, or insurance connections", "Prominent disclosure; never quietly change rankings"],
     ["B2B", "White-label or licensed workflows", "Separate buyer-facing disclosures and data boundaries"]],
    [1.1, 3.25, 2.45],
    8.4,
)

doc.add_page_break()

# Page 5
add_heading(doc, "Competitive landscape and wedge", 1)
p = doc.add_paragraph("The category is already moving beyond search. Several products combine inventory, finance, market data, AI guidance, or dealer outreach. The launch wedge should therefore be an independent decision workspace that can start with any listing or quote—not a claim to be the first AI car-shopping tool.")

competitive_rows = [
    ["CoPilot", "AI-assisted vehicle discovery and comparison", "Large consumer shopping experience"],
    ["CarEdge", "Market data, AI buying agent, dealer outreach, and human concierge", "Transparent pricing and negotiation workflow"],
    ["Capital One Auto Navigator", "Inventory plus personalized finance prequalification", "Financing integration and participating-dealer network"],
    ["CarXprt", "AI search, incentives, offer handling, and negotiation", "Search-to-deal workflow"],
    ["GetCarWise / CarClever", "AI-assistant inventory research and deal scoring", "Works inside general AI assistants"],
    ["Generic AI assistants", "Flexible questions, planning, and document interpretation", "Distribution and familiarity, but not always current or automotive-specific"],
]
add_table(doc, ["Alternative", "Current proposition", "Strength to respect"], competitive_rows, [1.55, 3.25, 2.0], 8.3)

add_heading(doc, "Recommended differentiation", 2)
for item in [
    "Bring-your-own listing and quote: the shopper can start anywhere instead of entering through a captive marketplace.",
    "Decision provenance: every recommendation shows the user's priority, the supporting fact or estimate, and the unresolved question.",
    "Out-the-door normalization: compare dealer offers on the same basis and show what changed.",
    "Consumer-controlled contact: research remains useful without becoming a lead; dealer outreach is a separate opt-in action.",
    "Portable buying brief: export a concise record for a partner, mechanic, lender, salesperson, or future session.",
    "Dealer readiness feedback loop: aggregate only consented, de-identified issue patterns to inform Dealer Growth Hackers' audit and consulting work.",
]:
    add_bullet(doc, item)

add_heading(doc, "Early product test", 2)
add_table(
    doc,
    ["Hypothesis", "Minimum evidence"],
    [["Shoppers will upload or paste a real listing", "At least 35% of qualified sessions add one listing"],
     ["A sourced comparison is more valuable than generic chat", "At least 50% of listing users view or export the decision brief"],
     ["Quote review creates willingness to pay", "At least 8% of eligible free users begin a premium checkout test"],
     ["Trust improves dealer-introduction quality", "Introduced users show higher appointment or purchase intent than undifferentiated leads"]],
    [3.65, 3.15],
    8.6,
)

doc.add_page_break()

# Page 6
add_heading(doc, "Domain shortlist", 1)
p = doc.add_paragraph("Status was checked against Verisign's authoritative .com RDAP service on October 1, 2026. “No record observed” means the registry returned no domain record at that moment; it is not a reservation, price quote, or trademark clearance.")

domain_rows = [
    ["1", "BetterCarBuyer.com", "No record observed", "Best brand", "Outcome-led, memorable, credible; AI can remain in the descriptor"],
    ["2", "MyNextCarAI.com", "No record observed", "Best AI-forward", "Personal and immediately clear; slightly long"],
    ["3", "ChooseMyCarAI.com", "No record observed", "Strong CTA", "Clear action and category; works well in ads"],
    ["4", "FindYourCarAI.com", "No record observed", "Search-led", "Direct promise, but may imply inventory breadth"],
    ["5", "SmartCarBuyerAI.com", "No record observed", "SEO descriptive", "Very clear but longest option"],
    ["6", "GuidedCarBuying.com", "No record observed", "Trust-led", "Professional and flexible; less app-like"],
    ["7", "CarFitAI.com", "No record observed", "Short", "Good for matching; narrower than the full buying journey"],
]
add_table(doc, ["Rank", "Domain", "Registry signal", "Role", "Assessment"], domain_rows, [0.45, 1.55, 1.25, 1.15, 2.4], 7.8)

add_heading(doc, "Names screened out", 2)
add_table(
    doc,
    ["Name", "Reason not recommended"],
    [["CarBuyingAI.com, AskCarAI.com, ShopCarsAI.com", "Already registered at the time checked"],
     ["CarClarityAI.com / DriveClarityAI.com", "No registry record observed, but “Clarity AI” creates avoidable brand-confusion risk"],
     ["CarTruthAI.com", "No registry record observed, but CarTruth is already used in automotive inspection"],
     ["GetCarWiseAI.com", "No registry record observed, but GetCarWise is an active automotive product"]],
    [2.3, 4.5],
    8.5,
)

add_heading(doc, "Domain action", 2)
add_step(doc, "01", "Run clearance", "Search federal and state trademarks, app stores, social handles, corporate names, and phonetic variants before using any name publicly.")
add_step(doc, "02", "Register defensively", "If cleared, secure the primary .com plus the strongest AI-forward redirect. Do not announce the name until checkout is complete.")
add_step(doc, "03", "Use one canonical brand", "Choose one public name and redirect backups. Multiple live names will dilute search signals and user trust.")

doc.add_page_break()

# Page 7
add_heading(doc, "Launch plan and decisions", 1)
roadmap_rows = [
    ["0–30 days", "Validate", "Interview 12–15 active shoppers; prototype intake, comparison, and decision brief; test name and pricing; identify data partners."],
    ["31–75 days", "Build MVP", "Ship responsive web app, accounts, listing and quote intake, recommendation rubric, scenario calculator, citations, consent, and analytics."],
    ["76–105 days", "Private beta", "Recruit 50–100 shoppers; review every high-stakes answer; measure completion, evidence use, export, willingness to pay, and support load."],
    ["106+ days", "Expand", "Add selected inventory, history, recall, incentive, and dealer-contact integrations only after source quality and economics are proven."],
]
add_table(doc, ["Timing", "Stage", "Outcome"], roadmap_rows, [1.0, 1.15, 4.65], 8.6)

add_heading(doc, "Launch metrics", 2)
for item in [
    "Activation: completes needs intake and adds at least one model, listing, or dealer quote",
    "Decision depth: compares two or more options and reviews the open-risk section",
    "Trust use: opens at least one source or edits an assumption before exporting",
    "Value: exports or shares a decision brief, returns to the project, or starts premium checkout",
    "Commercial quality: qualified dealer or service introductions, not raw lead volume",
    "Safety: material error rate, stale-data rate, consent defects, and escalations per 100 sessions",
]:
    add_bullet(doc, item)

add_heading(doc, "Decisions needed from Dealer Growth Hackers", 2)
decision_rows = [
    ["Brand", "Which domains are already owned, and should this be independent or explicitly “by Dealer Growth Hackers”?"],
    ["Audience", "U.S. only at launch? New, used, lease, EV, or all?"],
    ["Revenue", "Subscription, one-time report, concierge, referrals, B2B licensing, or a staged combination?"],
    ["Neutrality", "Will compensated inventory or partners ever influence ranking, and if so, how will that be disclosed?"],
    ["Data", "Which inventory, vehicle specification, recall, history, incentives, finance, and ownership-cost sources are available?"],
    ["Action", "Will the app advise only, draft dealer messages, or communicate on the shopper's behalf?"],
]
add_table(doc, ["Decision", "Question to resolve"], decision_rows, [1.25, 5.55], 8.7)

add_heading(doc, "Sources reviewed", 2)
sources = [
    ("AnswerThePublic", "Export for “ai car buying,” United States, English, October 1, 2026"),
    ("CoPilot", "https://www.copilotsearch.com/ai/"),
    ("CarEdge", "https://caredge.com/"),
    ("Capital One Auto Navigator", "https://www.capitalone.com/cars/about/"),
    ("CarXprt", "https://www.carxprt.com/"),
    ("GetCarWise / CarClever", "https://getcarwise.app/"),
    ("Verisign RDAP", "https://rdap.verisign.com/com/v1/"),
]
for label, url in sources:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(label + "  "), size=8.6, bold=True)
    if url.startswith("http"):
        add_hyperlink(p, url.replace("https://", "").rstrip("/"), url)
    else:
        set_font(p.add_run(url), size=8.6)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
